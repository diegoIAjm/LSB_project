import re
from django.contrib.auth.hashers import make_password
from openpyxl import load_workbook
from ..models import Usuario, Rol, Estudiante, Docente

class ImportService:
    
    @staticmethod
    def validar_email(email):
        """Validar formato de email"""
        email_regex = r'^[^\s@]+@[^\s@]+\.[^\s@]+$'
        return re.match(email_regex, email) is not None
    
    @staticmethod
    def validar_ci(ci):
        """Validar CI (solo números, 5-15 dígitos)"""
        ci_regex = r'^[0-9]{5,15}$'
        return re.match(ci_regex, ci) is not None
    
    @staticmethod
    def validar_nombre_apellido(texto):
        """Validar nombre/apellido (letras, espacios, 2-50 caracteres)"""
        nombre_regex = r'^[A-Za-zÁÉÍÓÚáéíóúÑñ\s]{2,50}$'
        return re.match(nombre_regex, texto) is not None
    
    @staticmethod
    def importar_estudiantes(archivo_excel):
        """
        Importar estudiantes desde archivo Excel
        Columnas esperadas: nombre, apellido, email, ci, nivel_actual
        """
        resultados = {
            'total': 0,
            'creados': 0,
            'duplicados': [],
            'errores': [],
            'detalle': []
        }
        
        try:
            wb = load_workbook(archivo_excel)
            ws = wb.active
            
            # Obtener el rol Estudiante
            try:
                rol_estudiante = Rol.objects.get(nombre__iexact='estudiante')
            except Rol.DoesNotExist:
                resultados['errores'].append('No se encontró el rol "Estudiante" en la base de datos')
                return resultados
            
            # Iterar sobre las filas (empezar desde la fila 2, asumiendo que fila 1 es cabecera)
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not any(row):
                    continue
                
                # Extraer datos (nombre, apellido, email, ci, nivel_actual)
                nombre = str(row[0]).strip() if row[0] else ''
                apellido = str(row[1]).strip() if row[1] else ''
                email = str(row[2]).strip().lower() if row[2] else ''
                ci = str(row[3]).strip() if row[3] else ''
                nivel_actual = str(row[4]).strip() if len(row) > 4 and row[4] else ''
                
                registro_info = {
                    'fila': row_idx,
                    'nombre': nombre,
                    'apellido': apellido,
                    'email': email,
                    'ci': ci,
                    'nivel_actual': nivel_actual
                }
                
                # Validaciones
                errores_fila = []
                
                # Validar nombre
                if not nombre or not ImportService.validar_nombre_apellido(nombre):
                    errores_fila.append('Nombre inválido (solo letras, 2-50 caracteres)')
                
                # Validar apellido
                if not apellido or not ImportService.validar_nombre_apellido(apellido):
                    errores_fila.append('Apellido inválido (solo letras, 2-50 caracteres)')
                
                # Validar email
                if not email or not ImportService.validar_email(email):
                    errores_fila.append('Email inválido')
                
                # Validar CI
                if not ci or not ImportService.validar_ci(ci):
                    errores_fila.append('CI inválido (solo números, 5-15 dígitos)')
                
                # Verificar duplicados en BD
                if email and Usuario.objects.filter(email=email).exists():
                    errores_fila.append(f'Email {email} ya existe')
                
                if ci and Usuario.objects.filter(ci=ci).exists():
                    errores_fila.append(f'CI {ci} ya existe')
                
                if errores_fila:
                    resultados['errores'].append({
                        'fila': row_idx,
                        'errores': errores_fila,
                        'datos': registro_info
                    })
                    continue
                
                # Crear usuario
                try:
                    usuario = Usuario.objects.create(
                        nombre=nombre,
                        apellido=apellido,
                        email=email,
                        password=make_password(ci),  # Contraseña = CI
                        rol=rol_estudiante,
                        ci=ci,
                        estado='activo'
                    )
                    
                    # Crear registro de estudiante
                    Estudiante.objects.create(
                        usuario=usuario,
                        nivel_actual=nivel_actual if nivel_actual else 'Principiante'
                    )
                    
                    resultados['creados'] += 1
                    resultados['detalle'].append({
                        'fila': row_idx,
                        'nombre': nombre,
                        'apellido': apellido,
                        'email': email,
                        'ci': ci,
                        'nivel_actual': nivel_actual,
                        'status': 'creado'
                    })
                    
                except Exception as e:
                    resultados['errores'].append({
                        'fila': row_idx,
                        'errores': [f'Error al guardar: {str(e)}'],
                        'datos': registro_info
                    })
            
            resultados['total'] = resultados['creados'] + len(resultados['errores'])
            
        except Exception as e:
            resultados['errores'].append({
                'fila': 0,
                'errores': [f'Error al leer el archivo: {str(e)}'],
                'datos': {}
            })
        
        return resultados
    
    @staticmethod
    def importar_docentes(archivo_excel):
        """
        Importar docentes desde archivo Excel
        Columnas esperadas: nombre, apellido, email, ci, especialidad
        """
        resultados = {
            'total': 0,
            'creados': 0,
            'duplicados': [],
            'errores': [],
            'detalle': []
        }
        
        try:
            wb = load_workbook(archivo_excel)
            ws = wb.active
            
            # Obtener el rol Docente
            try:
                rol_docente = Rol.objects.get(nombre__iexact='docente')
            except Rol.DoesNotExist:
                resultados['errores'].append('No se encontró el rol "Docente" en la base de datos')
                return resultados
            
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not any(row):
                    continue
                
                nombre = str(row[0]).strip() if row[0] else ''
                apellido = str(row[1]).strip() if row[1] else ''
                email = str(row[2]).strip().lower() if row[2] else ''
                ci = str(row[3]).strip() if row[3] else ''
                especialidad = str(row[4]).strip() if len(row) > 4 and row[4] else ''
                
                registro_info = {
                    'fila': row_idx,
                    'nombre': nombre,
                    'apellido': apellido,
                    'email': email,
                    'ci': ci,
                    'especialidad': especialidad
                }
                
                errores_fila = []
                
                if not nombre or not ImportService.validar_nombre_apellido(nombre):
                    errores_fila.append('Nombre inválido (solo letras, 2-50 caracteres)')
                
                if not apellido or not ImportService.validar_nombre_apellido(apellido):
                    errores_fila.append('Apellido inválido (solo letras, 2-50 caracteres)')
                
                if not email or not ImportService.validar_email(email):
                    errores_fila.append('Email inválido')
                
                if not ci or not ImportService.validar_ci(ci):
                    errores_fila.append('CI inválido (solo números, 5-15 dígitos)')
                
                if email and Usuario.objects.filter(email=email).exists():
                    errores_fila.append(f'Email {email} ya existe')
                
                if ci and Usuario.objects.filter(ci=ci).exists():
                    errores_fila.append(f'CI {ci} ya existe')
                
                if errores_fila:
                    resultados['errores'].append({
                        'fila': row_idx,
                        'errores': errores_fila,
                        'datos': registro_info
                    })
                    continue
                
                try:
                    usuario = Usuario.objects.create(
                        nombre=nombre,
                        apellido=apellido,
                        email=email,
                        password=make_password(ci),
                        rol=rol_docente,
                        ci=ci,
                        estado='activo'
                    )
                    
                    Docente.objects.create(
                        usuario=usuario,
                        especialidad=especialidad if especialidad else 'General'
                    )
                    
                    resultados['creados'] += 1
                    resultados['detalle'].append({
                        'fila': row_idx,
                        'nombre': nombre,
                        'apellido': apellido,
                        'email': email,
                        'ci': ci,
                        'especialidad': especialidad,
                        'status': 'creado'
                    })
                    
                except Exception as e:
                    resultados['errores'].append({
                        'fila': row_idx,
                        'errores': [f'Error al guardar: {str(e)}'],
                        'datos': registro_info
                    })
            
            resultados['total'] = resultados['creados'] + len(resultados['errores'])
            
        except Exception as e:
            resultados['errores'].append({
                'fila': 0,
                'errores': [f'Error al leer el archivo: {str(e)}'],
                'datos': {}
            })
        
        return resultados